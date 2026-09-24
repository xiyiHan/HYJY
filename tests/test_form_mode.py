r"""模板填充模式的离线测试。

运行：
    .venv\Scripts\python.exe tests\test_form_mode.py

覆盖：
  - 文字稿 Markdown 反解析（load_transcript）
  - 模型 JSON 回复的容错解析（含代码围栏、前后夹带说明的情况）
  - 结构化数据规整（缺字段、类型不对、会议类型非法）
  - 真实模板填充：以单位模板为底稿写入数据，并逐项回读验证

模板路径按以下顺序查找，找不到时相关用例跳过而不报错：
  1. 环境变量 MM_TEST_TEMPLATE
  2. config.yaml 的 template.path
  3. .cache/tplwork/会议纪要模板.docx
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass

from mmtools.config import Config, ConfigError  # noqa: E402
from mmtools.pipeline import load_transcript  # noqa: E402
from mmtools.summarizer import LLMError, normalize_form_data, parse_json_reply  # noqa: E402

PASSED: list[str] = []
FAILED: list[tuple[str, str]] = []
SKIPPED: list[str] = []


def check(name: str):
    def decorator(fn):
        try:
            fn()
            PASSED.append(name)
        except Exception as exc:  # noqa: BLE001
            FAILED.append((name, f"{type(exc).__name__}: {exc}"))
        return fn

    return decorator


def _find_template() -> Path | None:
    env = os.environ.get("MM_TEST_TEMPLATE", "").strip()
    if env and Path(env).exists():
        return Path(env)
    try:
        cfg = Config.load()
        configured = (cfg.get("template", "path") or "").strip()
        if configured and Path(configured).exists():
            return Path(configured)
    except Exception:
        pass
    fallback = ROOT / ".cache" / "tplwork" / "会议纪要模板.docx"
    if fallback.exists():
        return fallback
    return None


# ---------------------------------------------------------------- 文字稿反解析


@check("文字稿反解析：段落、时间戳、说话人均还原")
def _():
    cfg = Config.load()
    tmp = cfg.path("cache_dir") / "selftest" / "roundtrip.transcript.md"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        tmp.write_text(
            "# 转写文字稿\n\n"
            "- 音频文件：a.wav\n"
            "- 转写引擎：funasr / paraformer-zh\n"
            "- 说话人分离：已启用，识别到 2 人\n\n"
            "---\n\n"
            "**[00:01]** **说话人1**\n\n"
            "大家好，现在开会。\n\n"
            "**[01:05]** **说话人2**\n\n"
            "我先汇报进度。\n",
            encoding="utf-8",
        )
        result = load_transcript(tmp)
        assert result.backend == "funasr", f"引擎解析错：{result.backend}"
        assert result.model == "paraformer-zh"
        assert len(result.segments) == 2, f"段数错：{len(result.segments)}"
        assert result.segments[0].text == "大家好，现在开会。"
        assert result.segments[0].speaker == "说话人1"
        assert result.segments[0].start == 1.0
        assert result.segments[1].start == 65.0, f"时间戳错：{result.segments[1].start}"
        assert result.speakers == ["说话人1", "说话人2"]
    finally:
        tmp.unlink(missing_ok=True)


@check("文字稿反解析：单行也能还原且不串行")
def _():
    cfg = Config.load()
    tmp = cfg.path("cache_dir") / "selftest" / "single.transcript.md"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        tmp.write_text("**[00:05]** **说话人1**\n\n只有一句话。\n", encoding="utf-8")
        result = load_transcript(tmp)
        assert len(result.segments) == 1
        assert result.segments[0].text == "只有一句话。"
    finally:
        tmp.unlink(missing_ok=True)


# ---------------------------------------------------------------- JSON 解析


@check("JSON 解析：纯 JSON")
def _():
    data = parse_json_reply('{"topic": "监理例会", "items": ["a"]}')
    assert data["topic"] == "监理例会"
    assert data["items"] == ["a"]


@check("JSON 解析：带 Markdown 代码围栏")
def _():
    data = parse_json_reply('```json\n{"topic": "x"}\n```')
    assert data["topic"] == "x", data


@check("JSON 解析：前后夹带说明文字")
def _():
    data = parse_json_reply('好的，以下是结果：\n{"topic": "y"}\n希望对你有帮助。')
    assert data["topic"] == "y", data


@check("JSON 解析：非法内容给出可读报错")
def _():
    try:
        parse_json_reply("这完全不是 JSON")
    except LLMError as exc:
        assert "JSON" in str(exc)
        return
    raise AssertionError("非法内容未抛出 LLMError")


# ---------------------------------------------------------------- 数据规整


@check("数据规整：缺字段补空且类型正确")
def _():
    out = normalize_form_data({"topic": "监理例会"})
    assert out["topic"] == "监理例会"
    for key in ("project", "intro", "host", "owner_unit"):
        assert out[key] == "", f"{key} 未补空"
    assert out["items"] == []
    assert out["attendees"] == []
    assert out["signatories"] == {}


@check("数据规整：items 非列表时被包装")
def _():
    out = normalize_form_data({"items": "只有一条"})
    assert out["items"] == ["只有一条"]
    out = normalize_form_data({"items": ["a", "", "  ", "b"]})
    assert out["items"] == ["a", "b"], out["items"]


@check("数据规整：非法会议类型降级为其他")
def _():
    assert normalize_form_data({"meeting_type": "例会"})["meeting_type"] == "例会"
    assert normalize_form_data({"meeting_type": "茶话会"})["meeting_type"] == "其他"
    assert normalize_form_data({"meeting_type": ""})["meeting_type"] == ""


@check("数据规整：参会人名单类型兼容")
def _():
    out = normalize_form_data(
        {"attendees": [{"unit": "建设单位", "people": "张三,李四"}, "承建单位"]}
    )
    assert out["attendees"][0]["people"] == ["张三", "李四"], out["attendees"][0]
    assert out["attendees"][1] == {"unit": "承建单位", "people": []}


# ---------------------------------------------------------------- 模板填充


@check("真实模板：填充后各字段正确落位")
def _():
    template = _find_template()
    if template is None:
        SKIPPED.append("真实模板：填充后各字段正确落位（未找到模板文件）")
        return

    from docx import Document

    from mmtools.docx_template import fill_official_minutes
    from mmtools.transcribers import Segment, TranscriptResult

    cfg = Config.load()
    out = cfg.path("cache_dir") / "selftest" / "filled.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    data = {
        "project": "某某市政务信息化建设项目",
        "meeting_no": "第3次",
        "meeting_type": "例会",
        "topic": "监理例会",
        "location": "项目部会议室",
        "time": "2026年1月14日 上午",
        "host": "张工",
        "recorder": "李工",
        "publish_date": "2026年1月14日",
        "owner_unit": "某某市卫生健康委员会",
        "builder_unit": "某某科技有限公司",
        "supervisor_unit": "某某咨询",
        "signatories": {
            "建设单位": {"rep": "王主任", "date": "2026年1月15日"},
            "承建单位": {"rep": "赵经理", "date": "2026年1月15日"},
            "监理单位": {"rep": "张工", "date": "2026年1月15日"},
        },
        "intro": "2026年1月14日上午，某某市卫生健康委员会组织召开了某某市政务信息化建设项目第3次监理例会，会议主要围绕上阶段建设进度与存在问题展开讨论。",
        "items": ["承建单位于1月20日前提交设备到货计划。", "监理单位本周内完成隐蔽工程影像资料复核。", "设计单位1月18日前出具弱电井变更单。"],
        "attendees": [
            {"unit": "建设单位", "people": ["王主任"]},
            {"unit": "承建单位", "people": ["赵经理", "孙工"]},
            {"unit": "监理单位", "people": ["张工"]},
        ],
    }

    try:
        fill_official_minutes(cfg, template, out, data)
        assert out.exists() and out.stat().st_size > 0, "输出文件为空"

        doc = Document(str(out))

        # 表格：按标签回读
        table = doc.tables[0]

        def unique(cells):
            out = []
            for c in cells:
                if not out or out[-1]._tc is not c._tc:
                    out.append(c)
            return out

        def row_value(label: str) -> str:
            """回读标签在首列的字段值。"""
            for row in table.rows:
                cells = unique(row.cells)
                if cells and cells[0].text.strip().replace(" ", "") == label:
                    return cells[1].text.strip() if len(cells) > 1 else ""
            return "<未找到行>"

        def inline_value(label: str) -> str:
            """回读标签位于行中间的字段值。"""
            for row in table.rows:
                cells = unique(row.cells)
                for i, c in enumerate(cells):
                    if c.text.strip().replace(" ", "") == label and i + 1 < len(cells):
                        return cells[i + 1].text.strip()
            return "<未找到标签>"

        assert row_value("会议主题") == "监理例会", row_value("会议主题")
        assert row_value("会议地点") == "项目部会议室", row_value("会议地点")
        assert "2026年1月14日" in row_value("会议时间"), row_value("会议时间")
        assert row_value("会议记录") == "李工", row_value("会议记录")
        assert inline_value("会议主持") == "张工", inline_value("会议主持")
        assert "2026年1月14日" in inline_value("发布日期"), inline_value("发布日期")

        # 单位名称标签被替换为实际单位全称
        labels = [unique(r.cells)[0].text.strip() for r in table.rows]
        assert "某某市卫生健康委员会" in labels, f"建设单位名称未替换：{labels[8:12]}"
        assert "某某科技有限公司" in labels
        assert "某某咨询" in labels

        # 复选框：例会应被打勾（Wingdings 00FE）
        from docx.oxml.ns import qn

        type_cell = None
        for row in table.rows:
            if row.cells[0].text.strip() == "会议类型":
                type_cell = row.cells[1]
                break
        assert type_cell is not None, "未找到会议类型单元格"
        syms = type_cell._tc.findall(".//" + qn("w:sym"))
        chars = [s.get(qn("w:char")) for s in syms]
        assert "00FE" in chars, f"没有任何复选框被勾选：{chars}"
        assert chars.count("00FE") == 1, f"勾选了多个复选框：{chars}"

        # 正文
        body = "\n".join(p.text for p in doc.paragraphs)
        assert "某某市政务信息化建设项目" in body, "项目名称未写入"
        assert "（第3次）" in body, "会议次数未写入"
        assert "2026年1月14日上午" in body, "引言段未写入"
        assert "1、承建单位于1月20日前提交设备到货计划。" in body, "条目 1 未写入"
        assert "3、设计单位1月18日前出具弱电井变更单。" in body, "条目 3 未写入"
        assert "建设单位：王主任" in body, "出席名单未写入"
        assert "承建单位：赵经理  孙工" in body, f"多人姓名格式错：{body[-200:]}"
    finally:
        out.unlink(missing_ok=True)


@check("真实模板：条目多于模板槽位时自动增行")
def _():
    template = _find_template()
    if template is None:
        SKIPPED.append("真实模板：条目多于模板槽位时自动增行（未找到模板文件）")
        return

    from docx import Document

    from mmtools.docx_template import fill_official_minutes

    cfg = Config.load()
    out = cfg.path("cache_dir") / "selftest" / "filled_many.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        data = {
            "project": "测试项目",
            "meeting_no": "第1次",
            "meeting_type": "例会",
            "topic": "测试",
            "items": [f"第{i}条议定事项内容。" for i in range(1, 9)],
            "intro": "测试引言。",
            "attendees": [],
        }
        fill_official_minutes(cfg, template, out, data)
        doc = Document(str(out))
        body = "\n".join(p.text for p in doc.paragraphs)
        for i in range(1, 9):
            assert f"{i}、第{i}条议定事项内容。" in body, f"第 {i} 条未写入"
    finally:
        out.unlink(missing_ok=True)


@check("真实模板：原文件不被改动")
def _():
    template = _find_template()
    if template is None:
        SKIPPED.append("真实模板：原文件不被改动（未找到模板文件）")
        return

    from mmtools.docx_template import fill_official_minutes

    cfg = Config.load()
    before = template.read_bytes()
    out = cfg.path("cache_dir") / "selftest" / "copy_check.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        fill_official_minutes(cfg, template, out, {"project": "临时", "items": ["x"]})
        after = template.read_bytes()
        assert before == after, "模板原文件被修改了 —— 必须始终以副本为底稿"
    finally:
        out.unlink(missing_ok=True)


# ---------------------------------------------------------------- 主流程


def main() -> int:
    print("=" * 62)
    print("模板填充模式测试（不加载模型、不调用云端）")
    print("=" * 62)
    print("")

    for name in PASSED:
        print(f"  [通过] {name}")
    for name in SKIPPED:
        print(f"  [跳过] {name}")
    for name, err in FAILED:
        print(f"  [失败] {name}")
        print(f"         {err}")

    print("")
    print("-" * 62)
    total = len(PASSED) + len(FAILED)
    print(f"合计 {total} 项：通过 {len(PASSED)} 项，失败 {len(FAILED)} 项，跳过 {len(SKIPPED)} 项")
    if FAILED:
        return 1
    print("模板填充逻辑正常。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
