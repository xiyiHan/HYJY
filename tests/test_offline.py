r"""离线自测 —— 不加载模型、不调用云端，验证核心逻辑是否正确。

运行：
    .venv\Scripts\python.exe tests\test_offline.py

用途：
  1. 安装依赖前后都可以跑，用来确认代码与环境本身没问题；
  2. 转写或纪要报错时，先跑这个排除是代码问题还是模型/网络问题。
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from mmtools import prompts  # noqa: E402
from mmtools.config import Config, ConfigError  # noqa: E402
from mmtools.pipeline import _strip_markdown_scaffold  # noqa: E402
from mmtools.summarizer import CloudLLM, LLMError  # noqa: E402
from mmtools.transcribers import (  # noqa: E402
    Segment,
    TranscriptResult,
    format_clock,
    get_transcriber,
)

PASSED: list[str] = []
FAILED: list[tuple[str, str]] = []


def check(name: str):
    def decorator(fn):
        try:
            fn()
            PASSED.append(name)
        except Exception as exc:  # noqa: BLE001
            FAILED.append((name, f"{type(exc).__name__}: {exc}"))
        return fn

    return decorator


# ---------------------------------------------------------------- 配置


@check("配置加载与路径占位符解析")
def _():
    cfg = Config.load()
    workspace = str(cfg.get("paths", "workspace"))
    assert "{project_dir}" not in workspace, "占位符未被解析"
    assert "meeting-minutes" in workspace, f"工作目录解析异常：{workspace}"

    for key in ("workspace", "audio", "transcript", "minutes", "model_dir", "cache_dir"):
        # 走真实的 cfg.path() 接口，确保 paths 段定位正确
        value = str(cfg.path(key))
        assert "{workspace}" not in value and "{project_dir}" not in value
        # 本机 C 盘空间紧张，所有路径必须落在 D 盘
        assert value.upper().startswith("D:"), f"paths.{key} 未指向 D 盘：{value}"

    # 输出目录应逐层嵌套在 workspace 之下
    assert str(cfg.path("audio")).startswith(str(cfg.path("workspace")))


@check("缺失路径配置时报错可读")
def _():
    cfg = Config.load()
    try:
        cfg.path("not_exist_dir")
    except ConfigError as exc:
        assert "paths.not_exist_dir" in str(exc)
        return
    raise AssertionError("缺失路径时未抛出 ConfigError")


@check("默认平台为 deepseek 且含合规说明")
def _():
    cfg = Config.load()
    assert cfg.get("llm", "provider") == "deepseek"
    key, prov = cfg.provider()
    assert key == "deepseek"
    assert prov["base_url"].startswith("https://api.deepseek.com")
    assert prov.get("compliance"), "平台缺少 compliance 说明"


@check("平台注册表完整")
def _():
    cfg = Config.load()
    providers = cfg.providers()
    for expected in ("deepseek", "volcengine", "aliyun_bailian", "zhipu", "custom"):
        assert expected in providers, f"缺少平台 {expected}"
    assert len(providers) >= 5


@check("未知平台给出可读报错")
def _():
    cfg = Config.load()
    try:
        cfg.provider("not_exist_platform")
    except ConfigError as exc:
        assert "可用平台" in str(exc)
        return
    raise AssertionError("未对未知平台抛出 ConfigError")


@check("未配置密钥时的报错包含配置指引")
def _():
    cfg = Config.load()
    if cfg.api_key() is not None:
        return  # 已配置密钥，跳过
    try:
        cfg.require_api_key()
    except ConfigError as exc:
        text = str(exc)
        assert "环境变量" in text and "config.local.yaml" in text
        return
    raise AssertionError("未配置密钥时未抛出 ConfigError")


@check("密钥脱敏")
def _():
    cfg = Config.load()
    assert cfg.mask_secret("sk-1234567890abcdef") == "sk-123****cdef"
    assert cfg.mask_secret("short") == "sh****"
    assert cfg.mask_secret(None) == "(未配置)"


# ---------------------------------------------------------------- 转写数据结构


@check("时间戳格式化")
def _():
    assert format_clock(0) == "00:00"
    assert format_clock(65) == "01:05"
    assert format_clock(3661) == "01:01:01"
    assert format_clock(None) == "--:--"


@check("音频时长探测")
def _():
    import tempfile
    import wave

    from mmtools.transcribers import probe_duration

    cfg = Config.load()
    tmpdir = cfg.path("cache_dir") / "selftest"
    tmpdir.mkdir(parents=True, exist_ok=True)
    target = tmpdir / "probe.wav"
    try:
        with wave.open(str(target), "wb") as fh:
            fh.setnchannels(1)
            fh.setsampwidth(2)
            fh.setframerate(16000)
            fh.writeframes(b"\x00\x00" * 16000 * 3)  # 3 秒静音
        duration = probe_duration(target)
        assert duration is not None, "未能探测到时长为 None"
        assert abs(duration - 3.0) < 0.1, f"时长探测偏差：{duration}"
    finally:
        target.unlink(missing_ok=True)

    # 不存在的文件应返回 None 而不是抛异常
    assert probe_duration(tmpdir / "no_such.wav") is None


@check("转写结果转 Markdown")
def _():
    result = TranscriptResult(
        segments=[
            Segment(text="大家好，现在开会。", start=5, end=10, speaker="说话人1"),
            Segment(text="我先汇报进度。", start=12, end=18, speaker="说话人2"),
        ],
        backend="funasr",
        model="paraformer-zh",
        diarized=True,
        elapsed=120.0,
        duration=3600.0,
    )
    md = result.to_markdown("test.wav")
    assert "# 转写文字稿" in md
    assert "说话人1" in md and "说话人2" in md
    assert "识别到 2 人" in md
    assert "funasr / paraformer-zh" in md


@check("转写结果转提示词文本")
def _():
    result = TranscriptResult(
        segments=[Segment(text="附件里是初验方案。", start=65, end=70, speaker="说话人3")],
        backend="funasr",
        model="paraformer-zh",
    )
    text = result.to_prompt_text()
    assert text == "[01:05]说话人3 附件里是初验方案。", f"实际输出：{text}"


@check("说话人清单去重且保持顺序")
def _():
    result = TranscriptResult(
        segments=[
            Segment(text="a", speaker="说话人2"),
            Segment(text="b", speaker="说话人1"),
            Segment(text="c", speaker="说话人2"),
        ]
    )
    assert result.speakers == ["说话人2", "说话人1"]


@check("转写引擎工厂可实例化")
def _():
    cfg = Config.load()
    tr = get_transcriber(cfg)
    assert tr.name == "funasr"
    assert tr.conf.get("spk_model") == "cam++", "未启用说话人分离"
    tr2 = get_transcriber(cfg, "faster_whisper")
    assert tr2.name == "faster_whisper"


# ---------------------------------------------------------------- 长文切分


@check("短文本不切分")
def _():
    assert prompts.split_for_chunks("短文本", 1000) == ["短文本"]


@check("长文本切分且每块不超限")
def _():
    text = "\n".join(f"[00:{i:02d}]说话人1 这是第 {i} 行发言内容，用于测试切分逻辑。" for i in range(500))
    chunks = prompts.split_for_chunks(text, 2000)
    assert len(chunks) > 1, "长文本未被切分"
    for i, chunk in enumerate(chunks):
        assert len(chunk) <= 2000 + 200, f"第 {i} 块超长：{len(chunk)}"
    joined = "\n".join(chunks)
    assert "第 0 行" in joined and "第 499 行" in joined, "切分后内容丢失"


@check("切分保留重叠上下文")
def _():
    text = "\n".join(f"行{i}" for i in range(2000))
    chunks = prompts.split_for_chunks(text, 500, overlap_lines=2)
    assert len(chunks) > 1


# ---------------------------------------------------------------- 提示词


@check("纪要模板可取用")
def _():
    sup = prompts.get_system_prompt("supervision")
    gen = prompts.get_system_prompt("generic")
    assert "会议纪要" in sup and "议定事项" in sup
    assert "会议纪要" in gen
    # 反编造约束是纪要质量的关键，必须存在
    assert "严禁" in sup or "不得" in sup
    assert "（文字稿未提及）" in sup


@check("未知模板报错可读")
def _():
    try:
        prompts.get_system_prompt("not_exist_template")
    except KeyError as exc:
        assert "内置模板" in str(exc)
        return
    raise AssertionError("未对未知模板抛出 KeyError")


@check("用户提示词包含文字稿与会议名称")
def _():
    text = prompts.build_user_prompt("文字稿内容", meeting_title="监理例会", extra_notes="参会：三方")
    assert "监理例会" in text
    assert "文字稿内容" in text
    assert "<转写文字稿>" in text and "</转写文字稿>" in text
    assert "三方" in text


# ---------------------------------------------------------------- 文字稿清洗


@check("剥离文字稿元信息头")
def _():
    raw = (
        "<!--\n"
        "  生成时间：2026-09-24 10:00:00\n"
        "  转写引擎：funasr / paraformer-zh\n"
        "-->\n"
        "\n"
        "# 转写文字稿\n"
        "\n"
        "- 音频文件：a.wav\n"
        "- 转写引擎：funasr / paraformer-zh\n"
        "\n"
        "---\n"
        "\n"
        "**[00:05]** **说话人1**\n"
        "\n"
        "大家好，现在开会。\n"
    )
    body = _strip_markdown_scaffold(raw)
    assert "生成时间" not in body, "元信息未被剥离"
    assert "转写引擎" not in body, "元信息行未被剥离"
    assert "大家好，现在开会。" in body, "正文丢失"
    assert "说话人1" in body, "说话人行丢失"


@check("从已有文字稿元信息读取引擎名")
def _():
    from mmtools.pipeline import _backend_of_transcript

    cfg = Config.load()
    tmp = cfg.path("cache_dir") / "selftest" / "x.transcript.md"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        tmp.write_text(
            "# 转写文字稿\n\n- 转写引擎：funasr / paraformer-zh\n", encoding="utf-8"
        )
        assert _backend_of_transcript(tmp) == "funasr"

        tmp.write_text(
            "# 转写文字稿\n\n- 转写引擎：faster_whisper / small\n", encoding="utf-8"
        )
        assert _backend_of_transcript(tmp) == "faster_whisper"

        tmp.write_text("# 没有元信息\n正文\n", encoding="utf-8")
        assert _backend_of_transcript(tmp) == ""
    finally:
        tmp.unlink(missing_ok=True)


# ---------------------------------------------------------------- 术语规范化


@check("术语表：长词优先，避免被短词截断")
def _():
    from mmtools.glossary import build_rules

    rules = build_rules({"enabled": True, "replacements": {"监理": "X", "监理单位": "Y"}})
    assert len(rules) == 2
    # "监理单位" 必须排在 "监理" 之前，否则长词永远匹配不到
    assert rules[0].origin == "监理单位", f"排序错误：{[r.origin for r in rules]}"
    assert rules[1].origin == "监理"


@check("术语表：字面替换生效且保留时间戳与说话人")
def _():
    from mmtools.glossary import apply_to_result

    cfg = Config.load()
    cfg.data["transcription"]["glossary"] = {
        "enabled": True,
        "replacements": {"俊工": "竣工"},
    }
    result = TranscriptResult(
        segments=[
            Segment(text="项目已完成俊工验收。", start=10.0, end=15.0, speaker="说话人1"),
            Segment(text="下周补充资料。", start=20.0, end=25.0, speaker="说话人2"),
        ],
        backend="funasr",
        model="paraformer-zh",
    )
    updated, stats = apply_to_result(result, cfg)
    assert stats["enabled"] is True
    assert stats["count"] == 1, f"替换计数异常：{stats}"
    assert updated.segments[0].text == "项目已完成竣工验收。"
    # 时间戳与说话人不能被改动，否则无法与录音对照
    assert updated.segments[0].start == 10.0 and updated.segments[0].end == 15.0
    assert updated.segments[0].speaker == "说话人1"
    assert updated.segments[1].text == "下周补充资料。"


@check("术语表：未启用时原样返回")
def _():
    from mmtools.glossary import apply_to_result

    cfg = Config.load()
    cfg.data["transcription"]["glossary"] = {"enabled": False, "replacements": {"俊工": "竣工"}}
    result = TranscriptResult(segments=[Segment(text="俊工验收")], backend="funasr", model="x")
    updated, stats = apply_to_result(result, cfg)
    assert stats["enabled"] is False
    assert updated.segments[0].text == "俊工验收", "未启用却发生了替换"


@check("术语表：正则规则可用")
def _():
    from mmtools.glossary import apply_to_result, build_rules

    rules = build_rules(
        {"enabled": True, "regex": [{"pattern": "监理(工程师)?通知单", "replace": "监理通知单"}]}
    )
    assert len(rules) == 1 and rules[0].is_regex

    cfg = Config.load()
    cfg.data["transcription"]["glossary"] = {
        "enabled": True,
        "regex": [{"pattern": "监理(工程师)?通知单", "replace": "监理通知单"}],
    }
    result = TranscriptResult(
        segments=[Segment(text="已下发监理工程师通知单和监理通知单。")],
        backend="funasr",
        model="x",
    )
    updated, stats = apply_to_result(result, cfg)
    assert updated.segments[0].text == "已下发监理通知单和监理通知单。", updated.segments[0].text
    assert stats["count"] == 2


@check("术语表：非法正则给出可读报错")
def _():
    from mmtools.glossary import build_rules

    try:
        build_rules({"enabled": True, "regex": [{"pattern": "(未闭合", "replace": "x"}]})
    except ConfigError as exc:
        assert "无法编译" in str(exc)
        return
    raise AssertionError("非法正则未抛出 ConfigError")


@check("术语表：统计描述可读")
def _():
    from mmtools.glossary import describe_stats

    assert describe_stats({"enabled": False}) == "未启用"
    assert "未命中" in describe_stats({"enabled": True, "count": 0, "rules": 3})
    assert "修正 5 处" in describe_stats({"enabled": True, "count": 5, "rules": 3})


@check("热词分隔符规范化")
def _():
    from mmtools.transcribers import normalize_hotword

    assert normalize_hotword("甲,乙、丙 丁;戊") == "甲 乙 丙 丁 戊"
    assert normalize_hotword(["甲", "乙", "甲"]) == "甲 乙", "未去重"
    assert normalize_hotword("") == ""
    assert normalize_hotword(None) == ""
    assert normalize_hotword("  甲  ") == "甲"


@check("术语表已接入转写主流程（而不只是可单独调用）")
def _():
    import wave

    from mmtools.transcribers import run_transcription

    cfg = Config.load()
    cfg.data["transcription"]["glossary"] = {
        "enabled": True,
        "replacements": {"打磨院": "达摩院"},
    }

    class FakeTranscriber:
        """替身转写器，直接返回预设文本，用于验证后处理是否被调用。"""

        name = "fake"
        conf: dict = {}

        def transcribe(self, path):
            return TranscriptResult(
                segments=[Segment(text="欢迎大家体验打磨院的模型。", start=1.0, end=3.0)],
                backend="fake",
                model="fake",
            )

        def release(self):
            pass

    tmpdir = cfg.path("cache_dir") / "selftest"
    tmpdir.mkdir(parents=True, exist_ok=True)
    audio = tmpdir / "glossary_probe.wav"
    with wave.open(str(audio), "wb") as fh:
        fh.setnchannels(1)
        fh.setsampwidth(2)
        fh.setframerate(16000)
        fh.writeframes(b"\x00\x00" * 16000)

    try:
        result = run_transcription(cfg, audio, transcriber=FakeTranscriber())
        assert "达摩院" in result.segments[0].text, (
            f"术语表未在主流程中生效：{result.segments[0].text}"
        )
        assert "打磨院" not in result.segments[0].text
        # 元信息应带上术语表统计，便于追溯
        assert result.meta.get("glossary", {}).get("count") == 1
    finally:
        audio.unlink(missing_ok=True)


# ---------------------------------------------------------------- 导出


@check("Word 导出依赖就绪并生成有效文档")
def _():
    from mmtools.exporter import is_available, markdown_to_docx

    ok, note = is_available()
    if not ok:
        raise AssertionError(f"Word 导出依赖不可用：{note}")

    cfg = Config.load()
    tmp = cfg.path("cache_dir") / "selftest" / "export.docx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        md = (
            "# 会议纪要\n\n"
            "## 一、会议基本信息\n\n"
            "| 项目 | 内容 |\n| --- | --- |\n| 会议名称 | 监理例会 |\n\n"
            "## 二、议定事项\n\n"
            "1. **下周三前**报送方案\n"
        )
        markdown_to_docx(cfg, md, tmp)
        assert tmp.exists() and tmp.stat().st_size > 0, "Word 文档为空"

        from docx import Document

        doc = Document(str(tmp))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert any("会议纪要" in t for t in texts), "缺少标题"
        assert any("议定事项" in t for t in texts), "缺少正文"
        assert len(doc.tables) == 1, f"表格数量异常：{len(doc.tables)}"

        # 中文字体必须写到 w:eastAsia，只设 font.name 对中文不生效
        from docx.oxml.ns import qn

        found = False
        for p in doc.paragraphs:
            for r in p.runs:
                rpr = r._element.find(qn("w:rPr"))
                if rpr is None:
                    continue
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None and rfonts.get(qn("w:eastAsia")):
                    found = True
        assert found, "未设置东亚字体属性，中文在 Word 中会退回默认字体"
    finally:
        tmp.unlink(missing_ok=True)


# ---------------------------------------------------------------- 云端客户端


@check("云端客户端可构造且字段正确")
def _():
    cfg = Config.load()
    llm = CloudLLM(cfg)
    assert llm.endpoint == "https://api.deepseek.com/v1/chat/completions"
    assert llm.model == "deepseek-chat"
    assert "DeepSeek" in llm.label
    assert llm.compliance


@check("切换平台改变端点")
def _():
    cfg = Config.load()
    llm = CloudLLM(cfg, "aliyun_bailian")
    assert "dashscope.aliyuncs.com" in llm.endpoint
    assert "qwen" in llm.model.lower()


@check("自定义平台缺少 base_url 时报错可读")
def _():
    cfg = Config.load()
    try:
        CloudLLM(cfg, "custom")
    except ConfigError as exc:
        assert "base_url" in str(exc)
        return
    raise AssertionError("custom 平台未配置 base_url 时未报错")


@check("未配置密钥时调用失败并给出指引")
def _():
    cfg = Config.load()
    if cfg.api_key() is not None:
        return
    llm = CloudLLM(cfg)
    try:
        llm.chat([{"role": "user", "content": "hi"}])
    except ConfigError as exc:
        assert "API Key" in str(exc)
        return
    raise AssertionError("未配置密钥时未抛出 ConfigError")


@check("HTTP 错误分类给出对应建议")
def _():
    cfg = Config.load()
    llm = CloudLLM(cfg)

    class FakeResp:
        def __init__(self, code, text):
            self.status_code = code
            self.text = text

    cases = {
        401: "认证失败",
        404: "模型名",
        429: "限流",
        400: "请求被拒绝",
        500: "服务端异常",
    }
    for code, expect in cases.items():
        msg = llm._explain_http_error(FakeResp(code, "err"))
        assert expect in msg, f"HTTP {code} 的提示缺少关键词「{expect}」：{msg}"


# ---------------------------------------------------------------- 主流程


def main() -> int:
    print("=" * 62)
    print("离线自测（不加载模型，不调用云端）")
    print("=" * 62)
    print("")

    for name in PASSED:
        print(f"  [通过] {name}")
    for name, err in FAILED:
        print(f"  [失败] {name}")
        print(f"         {err}")

    print("")
    print("-" * 62)
    total = len(PASSED) + len(FAILED)
    print(f"合计 {total} 项：通过 {len(PASSED)} 项，失败 {len(FAILED)} 项")

    if FAILED:
        print("")
        print("存在失败项，请先修复后再进行转写与纪要生成。")
        return 1
    print("核心逻辑正常。")
    return 0


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    raise SystemExit(main())
