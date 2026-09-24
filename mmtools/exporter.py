"""会议纪要导出为 Word 文档。

为什么不用 pandoc：pandoc 生成的是通用排版，无法控制中文字体，导出的公文
既不是仿宋也不是黑体，直接归档不合格。这里用 python-docx 直接生成，
可以精确设置东亚字体（w:eastAsia 属性），产出符合中文公文习惯的文档。

依赖：python-docx（纯 Python，安装体积很小）。
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path

from .config import Config

# ---------------------------------------------------------------- 解析


_COMMENT_RE = re.compile(r"<!--(.*?)-->", re.S)
_TABLE_SEP_RE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\s*(\d+)[.、)]\s+(.*)$")


def _strip_comments(text: str) -> tuple[str, list[str]]:
    """剥离 HTML 注释，并返回其中的元信息行。"""
    meta: list[str] = []

    def _capture(match: re.Match) -> str:
        for line in match.group(1).splitlines():
            line = line.strip()
            if line:
                meta.append(line)
        return ""

    return _COMMENT_RE.sub(_capture, text), meta


def _split_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _iter_blocks(body: str):
    """把 Markdown 正文切成块：标题 / 表格 / 列表项 / 段落。"""
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line:
            i += 1
            continue

        if line in ("---", "***", "___"):
            i += 1
            continue

        # 表格：首行含 |，次行是分隔行
        if "|" in line and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]):
            header = _split_row(line)
            rows: list[list[str]] = []
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append(_split_row(lines[j]))
                j += 1
            yield ("table", header, rows)
            i = j
            continue

        m = _HEADING_RE.match(line)
        if m:
            yield ("heading", len(m.group(1)), m.group(2).strip())
            i += 1
            continue

        m = _BULLET_RE.match(raw)
        if m:
            yield ("bullet", m.group(1).strip())
            i += 1
            continue

        m = _ORDERED_RE.match(raw)
        if m:
            yield ("ordered", m.group(1), m.group(2).strip())
            i += 1
            continue

        yield ("para", line)
        i += 1


# ---------------------------------------------------------------- 生成


def _set_cjk_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    """同时设置西文与东亚字体。

    只设 run.font.name 在 Word 里对中文不生效，必须显式写 w:eastAsia，
    否则中文会退回默认字体，公文的字体要求就无法满足。
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def _add_rich_text(paragraph, text: str, font: str, size: float, base_bold: bool = False) -> None:
    """写入文本，识别 **加粗** 标记。"""
    for idx, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not chunk:
            continue
        # re.split 带捕获组时，奇数位是加粗内容
        bold = base_bold or (idx % 2 == 1)
        run = paragraph.add_run(chunk)
        _set_cjk_font(run, font, size, bold)


def markdown_to_docx(
    cfg: Config,
    markdown_text: str,
    output_path: Path,
    doc_title: str = "",
) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    conf = cfg.get("output", "docx") or {}
    title_font = conf.get("title_font", "黑体")
    heading_font = conf.get("heading_font", "黑体")
    body_font = conf.get("body_font", "仿宋")
    body_size = float(conf.get("body_size_pt", 12))
    title_size = float(conf.get("title_size_pt", 18))
    heading_size = float(conf.get("heading_size_pt", 14))
    line_spacing = float(conf.get("line_spacing_pt", 22))

    body, meta = _strip_comments(markdown_text)

    doc = Document()

    # 全局默认字体也要设，否则空段落、表格等会退回 Calibri
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(body_size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    normal.element.rPr.rFonts.set(qn("w:ascii"), body_font)
    normal.element.rPr.rFonts.set(qn("w:hAnsi"), body_font)

    doc_title_written = False

    for block in _iter_blocks(body):
        kind = block[0]

        if kind == "heading":
            level, text = block[1], block[2]
            if level == 1:
                # 一级标题作为文档大标题，居中
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(16)
                _add_rich_text(p, doc_title or text, title_font, title_size, base_bold=False)
                doc_title_written = True
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = Pt(line_spacing)
                _add_rich_text(p, text, heading_font, heading_size, base_bold=True)
            continue

        if kind == "table":
            _, header, rows = block
            if not header:
                continue
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, cell_text in enumerate(header):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                _add_rich_text(cell.paragraphs[0], cell_text, heading_font, body_size, base_bold=True)
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    value = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    _add_rich_text(cells[idx].paragraphs[0], value, body_font, body_size)
            doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = Pt(line_spacing)
        pf.space_after = Pt(3)

        if kind == "bullet":
            pf.left_indent = Pt(24)
            _add_rich_text(p, "· " + block[1], body_font, body_size)
        elif kind == "ordered":
            pf.left_indent = Pt(24)
            _add_rich_text(p, f"{block[1]}. {block[2]}", body_font, body_size)
        else:
            text = block[1]
            # 首行缩进两字符，符合中文公文习惯
            if not text.startswith("（") and not text.startswith("("):
                pf.first_line_indent = Pt(body_size * 2)
            _add_rich_text(p, text, body_font, body_size)

    if not doc_title_written and doc_title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_rich_text(p, doc_title, title_font, title_size)

    # 末尾加生成说明，样式低调但保证可追溯
    doc.add_paragraph()
    note_lines = [
        f"生成时间：{_dt.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        *[m for m in meta if not m.startswith("注意")],
        "本纪要由会议录音自动生成，归档前须经人工复核。",
    ]
    for line in note_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        _set_cjk_font(run, body_font, 9)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def is_available() -> tuple[bool, str]:
    try:
        import docx  # type: ignore  # noqa: F401
    except ImportError:
        return False, "未安装 python-docx，执行：pip install python-docx"
    return True, "依赖就绪"
