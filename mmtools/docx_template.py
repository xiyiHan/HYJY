"""按单位正式模板生成会议纪要。

与 exporter.py 的区别：
  - exporter.py  从零生成一份新的 Word 文档，排版由我们决定
  - 本模块       以单位既有的模板文件为底稿填充，**完整保留模板原有格式**

对政府项目的归档文件而言，后者才是正确做法 —— 模板里的字体、表格线、
复选框、页边距都是单位标准，重新排版会失去合规性。

实现要点：
  - 表格按**标签驱动**填充：先按行首标签（"会议主题"等）定位，再写入同行
    右侧的合并单元格，不硬编码行列坐标，模板微调后仍可用
  - 复选框是 Wingdings 符号 w:char="00A8"（空框），勾选即改为 "00FE"（打勾框），
    比替换成 Unicode 字符更能保持与原模板一致
  - 替换文字时保留首个 run 的格式，删除其余 run，避免字体被重置为默认值
"""

from __future__ import annotations

import copy
import shutil
from pathlib import Path

from .config import ConfigError

# Wingdings 复选框字符
BOX_UNCHECKED = "00A8"  # □
BOX_CHECKED = "00FE"  # ☑

# 模型在信息缺失时会填写这个占位文本
PLACEHOLDER = "（文字稿未提及）"


def _is_unknown(value) -> bool:
    text = str(value or "").strip()
    return (not text) or PLACEHOLDER in text


# ---------------------------------------------------------------- 基础操作


def _set_paragraph_text(paragraph, text: str):
    """替换段落文字，保留首个 run 的格式。

    直接给 paragraph.text 赋值会清掉所有格式；这里只改首个 run 的文字、
    删掉多余 run，从而完整继承模板里的字体与字号。
    """
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return paragraph
    runs[0].text = text
    for extra in runs[1:]:
        extra._element.getparent().remove(extra._element)
    return paragraph


def _set_cell_text(cell, text: str):
    """写入单元格：替换首段文字，并删除多余段落。"""
    if not cell.paragraphs:
        cell.text = text
        return cell
    _set_paragraph_text(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    return cell


def _cell_label(cell) -> str:
    return (cell.text or "").strip().replace(" ", "")


def _unique_cells(row) -> list:
    """把一行里因合并而重复出现的单元格去重，返回逻辑上的单元格序列。

    python-docx 对合并单元格会在每个 grid 位置返回同一个 tc 的引用，
    不去重就无法判断"标签后面那一格"到底是哪一格。
    """
    out = []
    for cell in row.cells:
        if not out or out[-1]._tc is not cell._tc:
            out.append(cell)
    return out


def _find_row(table, label: str):
    """按**首列**标签定位表格行。"""
    wanted = label.strip().replace(" ", "")
    for row in table.rows:
        cells = _unique_cells(row)
        if cells and _cell_label(cells[0]) == wanted:
            return row
    return None


def _value_cell(row, hint: str = ""):
    """取标签行中，第一个标签右侧的单元格（适用于标签在首列的字段）。"""
    cells = _unique_cells(row)
    if len(cells) < 2:
        return row.cells[-1]
    if hint:
        wanted = hint.strip().replace(" ", "")
        for i, cell in enumerate(cells):
            if _cell_label(cell) == wanted and i + 1 < len(cells):
                return cells[i + 1]
    return cells[1]


def _find_inline_value_cell(table, label: str):
    """定位「位于行中间」的标签，返回其右侧单元格。

    模板中"会议主持""发布日期"与"会议时间""会议记录"同处一行，
    标签不在首列，因此不能复用 _find_row 的逻辑。
    """
    wanted = label.strip().replace(" ", "")
    for row in table.rows:
        cells = _unique_cells(row)
        for i, cell in enumerate(cells):
            if _cell_label(cell) == wanted and i + 1 < len(cells):
                return cells[i + 1]
    return None


def _check_box(cell, selected: str) -> None:
    """把指定选项前的空框改为打勾框。"""
    from docx.oxml.ns import qn

    wanted = (selected or "").strip()
    if not wanted:
        return

    runs = cell.paragraphs[0].runs if cell.paragraphs else []
    opts: list[tuple[str, object]] = []
    pending_sym = None
    for run in runs:
        sym = run._element.find(qn("w:sym"))
        if sym is not None:
            pending_sym = sym
            continue
        text = (run.text or "").strip()
        if text and pending_sym is not None:
            opts.append((text, pending_sym))
            pending_sym = None

    for text, sym in opts:
        if text == wanted:
            sym.set(qn("w:char"), BOX_CHECKED)
            return


def _insert_paragraph_after(paragraph, text: str):
    """在指定段落后复制插入一个新段落，继承其全部段落格式。"""
    from docx.text.paragraph import Paragraph

    new_p = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    _set_paragraph_text(created, text)
    return created


def _delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


# ---------------------------------------------------------------- 表格填充


def _fill_header_table(table, data: dict) -> None:
    """填充「成文信息 / 签收信息 / 变更信息」表单。"""

    def put(label: str, value) -> None:
        """填充标签在首列的字段。"""
        if value is None:
            return
        row = _find_row(table, label)
        if row is None:
            return
        _set_cell_text(_value_cell(row), str(value))

    def put_inline(label: str, value) -> None:
        """填充标签位于行中间的字段（会议主持、发布日期）。"""
        if value is None:
            return
        cell = _find_inline_value_cell(table, label)
        if cell is not None:
            _set_cell_text(cell, str(value))

    row = _find_row(table, "会议类型")
    if row is not None:
        _check_box(_value_cell(row), str(data.get("meeting_type") or ""))

    put("会议主题", data.get("topic"))
    put("会议地点", data.get("location"))
    put("会议时间", data.get("time"))
    put_inline("会议主持", data.get("host"))
    put("会议记录", data.get("recorder"))
    put_inline("发布日期", data.get("publish_date"))

    # 签收信息：把「XX单位名称」标签替换为实际单位名，右侧填代表人与日期。
    # 这几项本来就是人工签署时手写的，未识别到就留空，不要把「（文字稿未提及）」
    # 填进正式文件的签署栏 —— 那是噪声，会让人误以为需要删除的内容。
    sign_rows = {
        "建设单位名称": data.get("owner_unit"),
        "承建单位名称": data.get("builder_unit"),
        "监理单位名称": data.get("supervisor_unit"),
    }
    sign_detail = data.get("signatories") or {}
    for label, unit_name in sign_rows.items():
        row = _find_row(table, label)
        if row is None:
            continue
        cells = _unique_cells(row)
        if not _is_unknown(unit_name):
            _set_cell_text(cells[0], str(unit_name))
        detail = sign_detail.get(label.replace("名称", "")) or {}
        rep = detail.get("rep") or ""
        date = detail.get("date") or ""
        if len(cells) >= 2:
            _set_cell_text(cells[1], "" if _is_unknown(rep) else str(rep))
        if len(cells) >= 3:
            _set_cell_text(cells[2], "" if _is_unknown(date) else str(date))


# ---------------------------------------------------------------- 正文填充


def _fill_body(doc, data: dict) -> None:
    """填充正文：标题、引言、编号条目、出席名单。"""
    paragraphs = doc.paragraphs

    def find_index(predicate, start: int = 0) -> int:
        for i in range(start, len(paragraphs)):
            if predicate(paragraphs[i]):
                return i
        return -1

    # 标题区：项目名称与会议次数
    if data.get("project"):
        idx = find_index(lambda p: "XXX项目" in p.text)
        if idx >= 0:
            _set_paragraph_text(paragraphs[idx], str(data["project"]))

    if data.get("meeting_no"):
        idx = find_index(lambda p: "第X次" in p.text or "（第" in p.text)
        if idx >= 0:
            _set_paragraph_text(paragraphs[idx], f"（{data['meeting_no']}）")

    # 引言段：模板中以"202X年"开头
    if data.get("intro"):
        idx = find_index(lambda p: p.text.strip().startswith("202X年"))
        if idx >= 0:
            _set_paragraph_text(paragraphs[idx], str(data["intro"]))

    _fill_items(paragraphs, data)
    _fill_attendees(paragraphs, data)


def _fill_items(paragraphs, data: dict) -> None:
    """填充「会议纪要如下」与「出席」之间的编号条目。

    只在两个锚点之间寻找编号段，避免误伤文档其他位置的数字开头段落。
    """
    items = [str(x).strip() for x in (data.get("items") or []) if str(x).strip()]
    if not items:
        return

    start = -1
    end = len(paragraphs)
    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        if start < 0 and text.startswith("会议纪要如下"):
            start = i + 1
        elif start >= 0 and text.startswith("出席"):
            end = i
            break
    if start < 0:
        return

    def is_numbered(p) -> bool:
        text = p.text.strip()
        return len(text) >= 2 and text[0].isdigit() and text[1] == "、"

    slots = [p for p in paragraphs[start:end] if is_numbered(p)]
    if not slots:
        return

    # 先按顺序写入已有槽位
    for i, paragraph in enumerate(slots):
        if i < len(items):
            _set_paragraph_text(paragraph, f"{i + 1}、{items[i]}")
        else:
            _set_paragraph_text(paragraph, "")

    # 条目多于槽位时，在最后一个槽位后依次复制插入，继承其段落格式
    if len(items) > len(slots):
        anchor = slots[-1]
        for offset, extra in enumerate(items[len(slots):], start=len(slots) + 1):
            anchor = _insert_paragraph_after(anchor, f"{offset}、{extra}")


def _fill_attendees(paragraphs, data: dict) -> None:
    attendees = data.get("attendees") or []
    if not attendees:
        return

    anchor_idx = -1
    for i, p in enumerate(paragraphs):
        if p.text.strip().startswith("出席"):
            anchor_idx = i
            break
    if anchor_idx < 0:
        return

    # 模板中"出席："后已有若干「单位简称：姓名 姓名」行，先收集再改写/增补
    template_lines = []
    for p in paragraphs[anchor_idx + 1:]:
        text = p.text.strip()
        if not text:
            break
        if "：" in text or ":" in text:
            template_lines.append(p)
        else:
            break

    anchor = paragraphs[anchor_idx]
    for i, line in enumerate(template_lines):
        if i < len(attendees):
            item = attendees[i]
            names = "  ".join(str(n) for n in (item.get("people") or []))
            label = str(item.get("unit") or "")
            _set_paragraph_text(line, f"{label}：{names}" if names else f"{label}：")
            anchor = line
        else:
            _set_paragraph_text(line, "")

    if len(attendees) > len(template_lines):
        for item in attendees[len(template_lines):]:
            names = "  ".join(str(n) for n in (item.get("people") or []))
            label = str(item.get("unit") or "")
            anchor = _insert_paragraph_after(anchor, f"{label}：{names}" if names else f"{label}：")


# ---------------------------------------------------------------- 对外接口


def fill_official_minutes(
    cfg: Config,
    template_path: Path,
    output_path: Path,
    data: dict,
) -> Path:
    """用数据填充模板，输出到指定路径。保留模板全部格式。"""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ConfigError("未安装 python-docx，执行：pip install python-docx") from exc

    template_path = Path(template_path)
    output_path = Path(output_path)
    if not template_path.exists():
        raise FileNotFoundError(f"模板文件不存在：{template_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    doc = Document(str(output_path))
    if doc.tables:
        _fill_header_table(doc.tables[0], data)
    _fill_body(doc, data)
    doc.save(str(output_path))
    return output_path


def is_available() -> tuple[bool, str]:
    try:
        import docx  # type: ignore  # noqa: F401
    except ImportError:
        return False, "未安装 python-docx，执行：pip install python-docx"
    return True, "依赖就绪"
